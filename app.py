import streamlit as st
from rag_pipeline import create_qa_chain_from_docs
from table_extractor import extract_tables_as_documents

from docx import Document as DocxDocument
from langchain_core.documents import Document as LangchainDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

# ----------------------------------
# Page Config
# ----------------------------------
st.set_page_config(
    page_title="Document Q&A",
    page_icon="📄",
    layout="centered"
)

# ----------------------------------
# Light Theme CSS
# ----------------------------------
st.markdown("""
<style>

.stApp {
    background: linear-gradient(135deg, #f8fafc, #eef2ff);
    color: #111827;
}

h1 {
    text-align: center;
    color: #1e3a8a;
    font-weight: 800;
}

.block-container {
    padding: 2rem 2rem;
}

.stTextInput input {
    background: #ffffff;
    color: #111827;
    border: 2px solid #c7d2fe;
    border-radius: 10px;
    padding: 10px;
}

.stFileUploader {
    background: #ffffff;
    border: 2px dashed #93c5fd;
    border-radius: 12px;
    padding: 12px;
}

.stButton button {
    background: linear-gradient(45deg, #3b82f6, #6366f1);
    color: white;
    font-weight: 600;
    border-radius: 10px;
    padding: 10px 18px;
    border: none;
}

.answer-box {
    background: #ffffff;
    border-left: 5px solid #3b82f6;
    padding: 15px;
    border-radius: 10px;
    color: #111827;
    box-shadow: 0 2px 8px rgba(0,0,0,0.05);
}

</style>
""", unsafe_allow_html=True)

# ----------------------------------
# Header
# ----------------------------------
st.markdown(
    "<h1>📄 RAG Powered Document Q&A</h1>",
    unsafe_allow_html=True
)

st.markdown(
    "<p style='text-align:center;color:#475569;'>Ask questions from your document easily</p>",
    unsafe_allow_html=True
)

st.divider()


# ----------------------------------
# Load DOCX preserving headings
# ----------------------------------
def load_docx_with_headings(path):
    doc = DocxDocument(path)
    full_text = ""
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Preserve heading structure with extra newlines
            if para.style.name.startswith("Heading"):
                full_text += "\n\n" + text + "\n\n"
            else:
                full_text += text + "\n\n"
    return [LangchainDocument(
        page_content=full_text,
        metadata={"source": path}
    )]


# ----------------------------------
# Upload File
# ----------------------------------
uploaded_file = st.file_uploader(
    "📂 Upload DOCX Document",
    type=["docx"]
)

# ----------------------------------
# Build RAG
# ----------------------------------
if uploaded_file:

    if (
        "current_file" not in st.session_state
        or st.session_state.current_file != uploaded_file.name
    ):

        try:
            temp_path = uploaded_file.name
            cache_name = uploaded_file.name.replace(".docx", "").replace(" ", "_")

            with open(temp_path, "wb") as f:
                f.write(uploaded_file.getbuffer())

            # --------------------------
            # Load text (heading-aware)
            # --------------------------
            with st.spinner("📖 Loading document text..."):
                documents = load_docx_with_headings(temp_path)

                splitter = RecursiveCharacterTextSplitter(
                    chunk_size=2000,
                    chunk_overlap=600,  # large overlap keeps clause numbers with headings
                    separators=[
                        "\n\n\n",
                        "\n\n",
                        "\n",
                        ". ",
                        " ",
                        ""
                    ]
                )

                text_docs = splitter.split_documents(documents)

            # --------------------------
            # Extract tables
            # --------------------------
            with st.spinner("📊 Extracting tables..."):
                table_docs = extract_tables_as_documents(
                    temp_path,
                    max_rows_per_chunk=20
                )

            # --------------------------
            # Combine
            # --------------------------
            docs = text_docs + table_docs

            st.write(f"📄 Text Chunks: {len(text_docs)}")
            st.write(f"📊 Table Chunks: {len(table_docs)}")
            st.write(f"📦 Total Chunks: {len(docs)}")

            # --------------------------
            # Build QA Chain
            # --------------------------
            import os
            cache_path = f"faiss_cache_{cache_name}"
            if os.path.exists(cache_path):
                st.info("⚡ Found existing cache — loading instantly!")
            else:
                st.info(f"⏳ First time loading — embedding {len(docs)} chunks. This may take a few minutes...")

            with st.spinner("🔧 Building QA chain..."):
                qa_chain = create_qa_chain_from_docs(
                    docs,
                    cache_name=cache_name
                )

            st.session_state.qa_chain = qa_chain
            st.session_state.current_file = uploaded_file.name

            st.success("✅ Document loaded successfully!")

        except Exception as e:
            st.error(f"❌ Processing failed: {e}")
            st.stop()

# ----------------------------------
# Query
# ----------------------------------
st.markdown("### ❓ Ask your question")

query = st.text_input("Type your question here...")

# ----------------------------------
# Ask
# ----------------------------------
if st.button("🚀 Get Answer"):

    if not st.session_state.get("qa_chain"):
        st.warning("Please upload a document first.")

    elif not query.strip():
        st.warning("Please enter a question.")

    else:
        qa_chain = st.session_state.qa_chain

        try:
            st.markdown("### 🔍 Retrieved Chunks (Debug)")
            retrieved_docs = qa_chain.retriever.invoke(query)

            for i, doc in enumerate(retrieved_docs):
                with st.expander(f"Chunk {i+1}"):
                    st.write(doc.page_content)

        except Exception as e:
            st.error(f"Retriever Debug Error: {e}")

        try:
            with st.spinner("🤖 Generating answer..."):
                response = qa_chain.invoke({"query": query})

            answer = response["result"]

            st.markdown("### 🧠 Response")
            st.markdown(
                f"<div class='answer-box'><b>{answer}</b></div>",
                unsafe_allow_html=True
            )

        except Exception as e:
            st.error(f"Answer Generation Error: {e}")