from docx import Document
from docx.oxml.ns import qn
from langchain_core.documents import Document as LangchainDocument


def extract_tables_as_documents(docx_path, max_rows_per_chunk=20):

    doc = Document(docx_path)
    documents = []

    # -----------------------------
    # Map table index → title
    # Table titles are paragraphs
    # sitting just above the table
    # in the document body
    # -----------------------------
    body = doc.element.body
    elements = list(body)

    table_titles = {}
    table_count = 0

    for i, elem in enumerate(elements):
        if elem.tag == qn('w:tbl'):
            title = ""
            # Check previous 1-2 siblings for a title paragraph
            for offset in [1, 2]:
                if i - offset >= 0:
                    prev = elements[i - offset]
                    if prev.tag == qn('w:p'):
                        prev_text = "".join(
                            node.text for node in prev.iter()
                            if node.tag == qn('w:t') and node.text
                        ).strip()
                        if prev_text:
                            title = prev_text
                            break
            table_titles[table_count] = title
            table_count += 1

    # -----------------------------
    # Extract tables
    # -----------------------------
    for table_idx, table in enumerate(doc.tables):

        rows = []

        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            # Remove duplicate merged cells
            row_data = list(dict.fromkeys(row_data))
            if any(row_data):
                rows.append(row_data)

        if not rows:
            continue

        title = table_titles.get(table_idx, "")
        headers = rows[0]
        data_rows = rows[1:]

        # Table with only header row
        if not data_rows:
            table_text = ""
            if title:
                table_text += f"{title}\n"
            table_text += f"Table {table_idx + 1}\n"
            table_text += " | ".join(headers) + "\n"

            documents.append(
                LangchainDocument(
                    page_content=table_text,
                    metadata={
                        "source": "full_table",
                        "table_id": table_idx + 1,
                        "title": title
                    }
                )
            )
            continue

        # Chunk table rows in groups of max_rows_per_chunk
        for i in range(0, len(data_rows), max_rows_per_chunk):
            chunk_rows = data_rows[i:i + max_rows_per_chunk]

            table_text = ""

            # Always include title in every chunk for searchability
            if title:
                table_text += f"{title}\n"

            table_text += f"Table {table_idx + 1} (Rows {i + 1} to {i + len(chunk_rows)})\n"
            table_text += " | ".join(headers) + "\n"
            table_text += "-" * 40 + "\n"

            for row in chunk_rows:
                row_parts = []
                for header, value in zip(headers, row):
                    if value:
                        row_parts.append(f"{header}: {value}")
                if row_parts:
                    table_text += " | ".join(row_parts) + "\n"

            documents.append(
                LangchainDocument(
                    page_content=table_text,
                    metadata={
                        "source": "table_chunk",
                        "table_id": table_idx + 1,
                        "title": title,
                        "row_start": i + 1,
                        "row_end": i + len(chunk_rows)
                    }
                )
            )

    return documents